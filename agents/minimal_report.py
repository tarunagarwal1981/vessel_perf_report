import streamlit as st
import pandas as pd
from io import BytesIO
from docx import Document

class MinimalReportGenerator:
    def __init__(self):
        self.report_generated = False
    
    def run(self, vessel_data, selected_vessel, hull_agent, speed_agent):
        st.header("Performance Report Generator")
        
        # Create tabs for report view and download
        report_tab1, report_tab2 = st.tabs(["Report Preview", "Download Report"])
        
        with report_tab1:
            # Display basic report preview
            st.subheader("Report Preview")
            
            # Extract basic metrics without creating charts
            hull_condition, power_loss, fuel_savings = self._get_hull_metrics(vessel_data, hull_agent)
            ballast_avg, laden_avg = self._get_speed_metrics(vessel_data)
            
            # Display preview
            st.markdown("### Hull Performance")
            st.markdown(f"**Hull Condition:** {hull_condition}")
            st.markdown(f"**Power Loss:** {power_loss:.1f}%")
            st.markdown(f"**Potential Fuel Savings:** {fuel_savings:.1f} MT/D")
            
            st.markdown("### Speed & Consumption")
            st.markdown(f"**Average Ballast Consumption:** {ballast_avg:.1f} MT/day")
            st.markdown(f"**Average Laden Consumption:** {laden_avg:.1f} MT/day")
        
        with report_tab2:
            st.markdown("### Generate Report")
            
            # Simple inputs
            report_date = st.date_input("Report Date")
            analyst_name = st.text_input("Analyst Name", "Marine Performance Team")
            
            # Generate report button
            if st.button("Generate Basic Report"):
                try:
                    with st.spinner("Generating report..."):
                        # Extract metrics without charts
                        hull_condition, power_loss, fuel_savings = self._get_hull_metrics(vessel_data, hull_agent)
                        ballast_avg, laden_avg = self._get_speed_metrics(vessel_data)
                        
                        # Create a simple document
                        doc = Document()
                        
                        # Add title
                        doc.add_heading(f'Vessel Performance Report - {selected_vessel.upper()}', 0)
                        doc.add_paragraph(f'Date: {report_date}')
                        doc.add_paragraph(f'Prepared by: {analyst_name}')
                        
                        # Hull performance section
                        doc.add_heading('Hull Performance', 1)
                        doc.add_paragraph(f'Hull Condition: {hull_condition}')
                        doc.add_paragraph(f'Power Loss: {power_loss:.1f}%')
                        doc.add_paragraph(f'Potential Fuel Savings: {fuel_savings:.1f} MT/D')
                        
                        # Speed section
                        doc.add_heading('Speed & Consumption', 1)
                        doc.add_paragraph(f'Average Ballast Consumption: {ballast_avg:.1f} MT/day')
                        doc.add_paragraph(f'Average Laden Consumption: {laden_avg:.1f} MT/day')
                        
                        # Appendix
                        doc.add_heading('Appendix', 1)
                        doc.add_paragraph('Hull Condition Rating:')
                        doc.add_paragraph('- Good: Power Loss < 15%')
                        doc.add_paragraph('- Average: Power Loss 15-25%')
                        doc.add_paragraph('- Poor: Power Loss > 25%')
                        
                        # Convert to bytes
                        docx_file = BytesIO()
                        doc.save(docx_file)
                        docx_file.seek(0)
                        
                        # Create download button
                        self.report_generated = True
                        st.download_button(
                            label="Download Report",
                            data=docx_file.getvalue(),
                            file_name=f"{selected_vessel}_Report_{report_date}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )
                        
                        st.success("Report generated successfully!")
                except Exception as e:
                    st.error(f"Error: {str(e)}")
    
    def _get_hull_metrics(self, vessel_data, hull_agent):
        # Extract basic hull metrics without charts
        try:
            # Filter data
            filtered_data = []
            for entry in vessel_data:
                if 'hull_roughness_power_loss' in entry and entry['hull_roughness_power_loss'] is not None:
                    filtered_data.append(entry)
            
            if filtered_data:
                # Get latest power loss
                power_loss = filtered_data[-1].get('hull_roughness_power_loss', 0)
                
                # Get condition
                condition, _ = hull_agent.get_hull_condition(power_loss)
                
                # Calculate savings
                fuel_savings = (power_loss - 15) * 0.05 if power_loss > 15 else 0
                
                return condition, power_loss, fuel_savings
            else:
                return "Unknown", 0, 0
        except:
            return "Error", 0, 0
    
    def _get_speed_metrics(self, vessel_data):
        # Extract basic speed metrics without charts
        try:
            # Filter data
            ballast_consumptions = []
            laden_consumptions = []
            
            for entry in vessel_data:
                if 'loading_condition' in entry and 'normalised_consumption' in entry:
                    if entry['normalised_consumption'] is not None:
                        if entry.get('loading_condition', '').lower() == 'ballast':
                            ballast_consumptions.append(entry['normalised_consumption'])
                        elif entry.get('loading_condition', '').lower() == 'laden':
                            laden_consumptions.append(entry['normalised_consumption'])
            
            # Calculate averages
            ballast_avg = sum(ballast_consumptions) / len(ballast_consumptions) if ballast_consumptions else 0
            laden_avg = sum(laden_consumptions) / len(laden_consumptions) if laden_consumptions else 0
            
            return ballast_avg, laden_avg
        except:
            return 0, 0
